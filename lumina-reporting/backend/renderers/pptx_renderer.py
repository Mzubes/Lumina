from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt

def render_pptx(content):
    presentation = Presentation()

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = content['report_title']
    title_slide.placeholders[1].text = (
        f"{content.get('client_name') or ''}\nGenerated: {content.get('generated_at') or ''}"
    )

    blank_layout = presentation.slide_layouts[6]
    for component in content['components']:
        slide = presentation.slides.add_slide(blank_layout)

        title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(9), Inches(0.8))
        title_paragraph = title_box.text_frame.paragraphs[0]
        title_paragraph.text = component['title']
        title_paragraph.font.size = Pt(28)
        title_paragraph.font.bold = True

        if component['type'] == 'text_block':
            body_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(9), Inches(5))
            body_box.text_frame.word_wrap = True
            body_box.text_frame.text = component.get('text', '')
        elif component['type'] == 'people_grid':
            body_box = slide.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(9), Inches(5))
            body_box.text_frame.word_wrap = True
            for index, person in enumerate(component.get('people', [])):
                paragraph = body_box.text_frame.paragraphs[0] if index == 0 else body_box.text_frame.add_paragraph()
                subtitle = ' · '.join(part for part in [person.get('title'), person.get('detail')] if part)
                paragraph.text = f"{person.get('name', '')}" + (f" — {subtitle}" if subtitle else '')
        else:
            columns = component.get('columns', [])
            rows = component.get('rows', [])
            if columns:
                row_count = len(rows) + 1
                table_shape = slide.shapes.add_table(
                    row_count, len(columns), Inches(0.4), Inches(1.3), Inches(9), Inches(0.4 * row_count),
                )
                table = table_shape.table
                for col_index, column_name in enumerate(columns):
                    table.cell(0, col_index).text = str(column_name)
                for row_index, row in enumerate(rows, start=1):
                    for col_index, value in enumerate(row):
                        table.cell(row_index, col_index).text = str(value)

    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
